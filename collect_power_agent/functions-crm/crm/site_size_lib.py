"""site_size_lib.py -- measure site size (sitemap page count) for the prospect catalogue.

Pre-step of app/prospects_import.py (--measure):

  1. collect the unique sites in the catalogue rows whose `page_count` cell is empty
     (all of them with force=True),
  2. read each site's sitemap with crm.sitemap_reader.SitemapReader,
  3. write the result back to the spreadsheet columns `page_count` and `sitemap_url`
     (columns are appended after the last used header when missing),
  4. the import then picks the values up like any other column
     (campaign_leads.page_count / .sitemap_url).

Empty cells are only filled; existing values are replaced only with force=True.
A site without a readable sitemap is left empty and retried on the next run.
"""
from __future__ import annotations

import asyncio
import shutil
from pathlib import Path

from crm.prospect_import_lib import (
    _cell, _field_for, _website, lead_id_from_url)

SIZE_COLS = ("page_count", "sitemap_url")      # sheet column header == lead field
READ_TIMEOUT = 120.0


def site_key(row: dict) -> tuple[str, str]:
    """(lead_id, website) for a parsed row; ('', '') when it has no site."""
    web = _website(row.get("website", ""), row.get("contact_page", ""))
    return (lead_id_from_url(web), web) if web else ("", "")


def sites_to_measure(rows: list[dict], *, force: bool = False) -> dict[str, str]:
    """{lead_id: website} for rows that need a measurement."""
    todo: dict[str, str] = {}
    for r in rows:
        lid, web = site_key(r)
        if lid and (force or not r.get("page_count")):
            todo.setdefault(lid, web)
    return todo


def _say(msg: str) -> None:
    print(msg, flush=True)


def _fmt_eta(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 60}m{sec % 60:02d}s" if sec >= 60 else f"{sec}s"


async def _measure_async(sites: dict[str, str], workers: int, log) -> dict[str, dict]:
    import time
    import aiohttp
    from crm.sitemap_reader import SitemapReader
    total = len(sites)
    width = len(str(total))
    sem = asyncio.Semaphore(max(1, workers))
    results: dict[str, dict] = {}
    counts = {"done": 0, "none": 0, "timeout": 0, "error": 0}
    t0 = time.monotonic()
    headers = {"Accept-Language": "en,da;q=0.9"}
    async with aiohttp.ClientSession(
            connector=aiohttp.TCPConnector(limit=workers * 2, ssl=False),
            headers=headers) as session:

        async def one(lid: str, web: str) -> None:
            async with sem:
                t1 = time.monotonic()
                outcome, detail = "none", "no sitemap found"
                try:
                    pc, sm_url, *_ = await asyncio.wait_for(
                        SitemapReader(session, web).read(), timeout=READ_TIMEOUT)
                    if pc and pc > 0:
                        results[lid] = {"page_count": int(pc), "sitemap_url": sm_url or ""}
                        outcome = "ok"
                        detail = f"{int(pc):,} pages  {sm_url or ''}"
                except asyncio.TimeoutError:
                    outcome, detail = "timeout", f"no answer in {int(READ_TIMEOUT)}s"
                except Exception as exc:                      # never stop the run
                    outcome, detail = "error", f"{type(exc).__name__}: {str(exc)[:80]}"
                counts["done"] += 1
                if outcome != "ok":
                    counts[outcome] += 1
                elapsed = time.monotonic() - t0
                eta = elapsed / counts["done"] * (total - counts["done"])
                log(f"[measure] [{counts['done']:>{width}}/{total}] {outcome.upper():<7} "
                    f"{web}  {detail}  ({time.monotonic() - t1:.1f}s, "
                    f"{'done' if counts['done'] == total else 'ETA ' + _fmt_eta(eta)})")

        await asyncio.gather(*(one(l, w) for l, w in sites.items()))
    log(f"[measure] finished in {_fmt_eta(time.monotonic() - t0)}: {len(results)} measured, "
        f"{counts['none']} without sitemap, {counts['timeout']} timed out, "
        f"{counts['error']} failed (those stay empty and are retried next run)")
    return results


def measure_sites(sites: dict[str, str], *, workers: int = 8, log=_say) -> dict[str, dict]:
    """{lead_id: {'page_count': n, 'sitemap_url': url}} for sites with a readable sitemap."""
    if not sites:
        return {}
    log(f"[measure] reading {len(sites)} sitemap(s), {workers} at a time "
        f"(max {int(READ_TIMEOUT)}s each) ...")
    return asyncio.run(_measure_async(sites, workers, log))


def apply_to_rows(rows: list[dict], results: dict[str, dict], *, force: bool = False) -> int:
    """Patch parsed rows in memory (so the import sees values even in a dry run)."""
    n = 0
    for r in rows:
        res = results.get(site_key(r)[0])
        if not res:
            continue
        for col in SIZE_COLS:
            if res.get(col) not in ("", None) and (force or not r.get(col)):
                r[col] = str(res[col])
                n += 1
    return n


def write_back(paths: list[Path], results: dict[str, dict], *, force: bool = False,
               apply: bool = False, log=_say) -> dict[str, int]:
    """Write page_count / sitemap_url into the catalogue workbooks.
    Without apply nothing is saved.  A '.xlsx.bak' copy is kept next to each saved file."""
    from openpyxl import load_workbook
    stats: dict[str, int] = {}
    for path in paths:
        wb = load_workbook(str(path))                 # full mode keeps formatting
        changed = 0
        for ws in wb.worksheets:
            header = [c.value for c in ws[1]]
            fields = [_field_for(h) for h in header]
            if "email" not in fields or "company" not in fields:
                continue
            col_idx: dict[str, int] = {}
            for i, f in enumerate(fields, start=1):
                if f in SIZE_COLS and f not in col_idx:
                    col_idx[f] = i
            last = max((i for i, h in enumerate(header, start=1) if h not in (None, "")),
                       default=0)
            for col in SIZE_COLS:
                if col not in col_idx:
                    last += 1
                    col_idx[col] = last
                    ws.cell(row=1, column=last, value=col)
            for row in ws.iter_rows(min_row=2):
                rec = {}
                for f, c in zip(fields, row):
                    if f and _cell(c.value) and not rec.get(f):
                        rec[f] = _cell(c.value)
                res = results.get(site_key(rec)[0])
                if not res:
                    continue
                for col in SIZE_COLS:
                    cell = ws.cell(row=row[0].row, column=col_idx[col])
                    if res.get(col) in ("", None):
                        continue
                    if force or _cell(cell.value) == "":
                        if _cell(cell.value) != str(res[col]):
                            cell.value = res[col]
                            changed += 1
        stats[path.name] = changed
        if changed and apply:
            shutil.copy2(path, str(path) + ".bak")
            wb.save(str(path))
            log(f"[measure] {path.name}: {changed} cell(s) written (backup: {path.name}.bak)")
        else:
            log(f"[measure] {path.name}: {changed} cell(s) "
                f"{'would be written' if changed else 'to write'}")
        wb.close()
    return stats
