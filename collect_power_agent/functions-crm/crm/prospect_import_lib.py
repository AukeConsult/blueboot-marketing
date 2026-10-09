"""prospect_import_lib.py -- Sync the BlueSearch prospect catalogue (xlsx) into campaigns.

The catalogue is a folder of Excel files (one per country, plus an "international"
file).  Every sheet with 'Company' and 'Email' headers is read; rows are grouped
by COUNTRY, so each country gets its own campaign  <prefix>_<CC>  (e.g. BS_UK).
Rows from every file are merged, so the same company/contact appearing in two
files ends up once.

Writes (all idempotent -- run it as often as you like)
------------------------------------------------------
Always:
  campaigns/{campaign}                                 created as draft if missing
  campaigns/{campaign}/campaign_leads/{lead_id}        one per website
  campaigns/{campaign}/campaign_contacts/{doc_id}      one per email address
Unless campaign_only:
  site_leads/{lead_id}                                 created only when missing
  site_leads/{lead_id}/site_contacts/{contact_id}      created / empty fields filled
  email_contacts/{doc_id}                              created / empty fields filled

New vs changed
--------------
Each document is compared with what is already in Firestore:
  new        document does not exist yet -> created
  changed    one or more catalogue fields differ -> only those fields updated
  unchanged  nothing to do
A blank cell in the catalogue never blanks out a value in Firestore.
New campaign_leads get status "pending" -- the outreach send skips contacts whose lead
has no status.  Existing leads with an empty status are repaired the same way.
Outreach state is never touched on existing contacts: status, mail_sent,
followup_*, comment_history, sent_at, message_id ... (see PROTECTED_CONTACT_FIELDS).
Rows are never deleted; contacts that disappeared from the catalogue are reported.

IDs (same rules as the rest of the system)
------------------------------------------
  doc_id     re.sub(r'[^a-zA-Z0-9_-]', '_', email.lower())      (email_contacts rule)
  lead_id    host of the website, dots/hyphens -> '_'            (lead_id_from_url)
  contact_id sha1(email.lower())[:12]                            (site_agent rule)
"""
from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from crm.contact_clean_lib import find_emails

CAMPAIGNS_COLLECTION    = "campaigns"
CAMPAIGN_LEADS_SUB      = "campaign_leads"
CAMPAIGN_CONTACTS_SUB   = "campaign_contacts"
SITE_LEADS_COLLECTION   = "site_leads"
SITE_CONTACTS_SUB       = "site_contacts"
EMAIL_CONTACTS_COLLECTION = "email_contacts"

SOURCE      = "prospect_import"
DEFAULT_PREFIX = "BS"
BATCH_SIZE  = 400

# Existing-contact fields the import must never overwrite (outreach history).
PROTECTED_CONTACT_FIELDS = {
    "status", "mail_sent", "next_mail_index", "in_reply_to",
    "followup_status", "followup_date", "followup_comment",
    "followup_importance", "followup_owner", "comment_history",
    "sent_at", "message_id", "sender_account", "created_at", "added_at",
    "last_action", "last_action_status",
    "send_confirmation", "send_confirmed_by", "send_confirmed_at", "send_confirm_note",
}

COUNTRY_CODES = {
    "denmark": "DK", "finland": "FI", "germany": "DE", "sweden": "SE",
    "norway": "NO", "united kingdom": "UK", "uk": "UK", "great britain": "UK",
    "gb": "UK", "england": "UK", "united states": "US", "usa": "US", "us": "US",
    "netherlands": "NL", "france": "FR", "spain": "ES", "italy": "IT",
    "ireland": "IE", "poland": "PL", "canada": "CA", "australia": "AU",
}
COUNTRY_NAMES = {
    "DK": "Denmark", "FI": "Finland", "DE": "Germany", "SE": "Sweden",
    "NO": "Norway", "UK": "United Kingdom", "US": "United States",
    "NL": "Netherlands", "FR": "France", "ES": "Spain", "IT": "Italy",
    "IE": "Ireland", "PL": "Poland", "CA": "Canada", "AU": "Australia",
}

_HEADERS = {
    "priority": "priority", "company": "company", "country": "country",
    "city_state": "location", "prospect_type": "prospect_type",
    "contact_person": "name", "role_to_approach": "title", "email": "email",
    "email_type": "email_type_raw", "phone": "phone", "contact_url": "contact_page",
    "site": "website", "suggested_sales_angle": "suggested_angle",
    "company_or_person_source": "company_source", "contact_source": "contact_source",
    "source_urls": "source_urls", "research_date": "researched_at",
    "status": "prospect_status", "last_contacted": "prospect_last_contacted",
    "next_follow_up": "prospect_next_followup", "notes": "notes",
    # written by the --measure pre-step (site_size_lib); "site_size"/"sitesize"/"sitesie" accepted
    "page_count": "page_count", "site_size": "page_count", "sitesize": "page_count",
    "sitesie": "page_count", "sitemap_url": "sitemap_url",
}

_GENERIC_LOCALS = {
    "info", "hello", "hi", "contact", "kontakt", "post", "mail", "office", "sales",
    "support", "enquiries", "enquiry", "team", "admin", "service", "kundenservice",
    "firmapost", "kontor", "hej", "hallo", "moin", "ahoi", "hei", "asiakaspalvelu",
    "myynti", "business", "orders", "order", "accessibility", "marketing",
}


# ---------------------------------------------------------------------------
# ID helpers
# ---------------------------------------------------------------------------

def doc_id_from_email(email: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]", "_", email.lower().strip())


def contact_id_from_email(email: str) -> str:
    return hashlib.sha1(email.lower().encode()).hexdigest()[:12]


def lead_id_from_url(url: str) -> str:
    host = urlparse(url).hostname or url
    slug = re.sub(r"[.\-]+", "_", host.rstrip(".").lower())
    return re.sub(r"_+", "_", slug).strip("_")


def domain_from_url(url: str) -> str:
    host = (urlparse(url).hostname or "").lower()
    return host[4:] if host.startswith("www.") else host


# ---------------------------------------------------------------------------
# Reading the catalogue
# ---------------------------------------------------------------------------

def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").lower()).strip("_")


def _field_for(h) -> str | None:
    n = _norm_header(h)
    if n.startswith("why_"):
        return "description"
    return _HEADERS.get(n)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _website(site: str, contact_url: str) -> str:
    raw = (site or "").strip() or (contact_url or "").strip()
    if not raw:
        return ""
    if not re.match(r"^https?://", raw, re.I):
        raw = "https://" + raw
    p = urlparse(raw)
    if not p.hostname:
        return ""
    if site.strip():
        return raw.rstrip("/")
    return f"{p.scheme}://{p.hostname}"       # from contact URL: keep only the origin


def _priority(raw: str) -> tuple[str, int | None]:
    raw = (raw or "").strip()
    if re.fullmatch(r"\d+(\.0+)?", raw):
        return "High", int(float(raw))          # numbered rows rank above Medium
    return raw.capitalize(), None


def _email_type(raw: str, email: str) -> str:
    local = email.split("@")[0].lower() if email else ""
    r = (raw or "").lower()
    if local in _GENERIC_LOCALS:
        return "generic"
    if "named" in r or "personal" in r:
        return "personal"
    return "generic" if r else "personal"


def parse_catalogue(paths: list) -> tuple[list[dict], list[str]]:
    """Read all prospect sheets.  Returns (rows, warnings); rows carry '_src'.
    Each item is a Path, or a (name, xlsx_bytes) tuple for an uploaded file."""
    import io
    from openpyxl import load_workbook
    rows: list[dict] = []
    warnings: list[str] = []
    for path in paths:
        if isinstance(path, tuple):
            name, blob = path
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            path = Path(name)
        else:
            wb = load_workbook(str(path), read_only=True, data_only=True)
        found = False
        for ws in wb.worksheets:
            data = list(ws.iter_rows(values_only=True))
            if not data:
                continue
            fields = [_field_for(h) for h in data[0]]
            if "email" not in fields or "company" not in fields:
                continue
            found = True
            for n, raw in enumerate(data[1:], start=2):
                rec: dict = {}
                for f, v in zip(fields, raw):
                    if f and _cell(v) and not rec.get(f):
                        rec[f] = _cell(v)
                if not rec.get("company") and not rec.get("email"):
                    continue
                rec["_src"] = f"{path.name}:{ws.title}:{n}"
                rows.append(rec)
        if not found:
            warnings.append(f"{path.name}: no sheet with Company + Email headers -- skipped")
    return rows, warnings


# ---------------------------------------------------------------------------
# Normalising rows -> lead / contact dicts, grouped by campaign
# ---------------------------------------------------------------------------

def country_code(cname: str) -> str:
    """'United Kingdom' / 'UK' / 'US and Canada' -> 'UK' / 'UK' / 'US'; '' when unknown."""
    cname = (cname or "").strip()
    code = COUNTRY_CODES.get(cname.lower())
    if not code:        # "US and Canada", "Sweden / Norway" -> first country named
        first = re.split(r"\s+and\s+|[/,;&]", cname, maxsplit=1)[0].strip().lower()
        code = COUNTRY_CODES.get(first)
    if not code and re.fullmatch(r"[A-Za-z]{2}", cname):
        code = cname.upper()
    return code or ""


def _page_count(raw: str):
    digits = re.sub(r"[^\d]", "", str(raw or "").split(".")[0])
    return int(digits) if digits else None


def _merge_into(dst: dict, src: dict) -> None:
    for k, v in src.items():
        if v not in ("", None):
            dst[k] = v


def build_records(rows: list[dict], *, prefix: str = DEFAULT_PREFIX,
                  only_campaign: str = "", countries: set[str] | None = None
                  ) -> tuple[dict, list[str]]:
    """Returns ({campaign: {'leads': {lead_id: doc}, 'contacts': {doc_id: doc},
    'contacted_skipped': [...], 'no_email': n}}, warnings)."""
    warnings: list[str] = []
    out: dict[str, dict] = {}

    # Older research first, so newer rows win when the same record appears twice.
    ordered = sorted(rows, key=lambda r: r.get("researched_at", ""))
    no_email_rows: dict[str, int] = {}

    for r in ordered:
        src = r["_src"]
        cname = (r.get("country") or "").strip()
        code = country_code(cname)
        if not code:
            warnings.append(f"{src}: unknown country '{cname}' -- skipped")
            continue
        if countries and code not in countries:
            continue
        campaign = only_campaign or f"{prefix}_{code}"

        raw_email = r.get("email", "")
        found = find_emails(raw_email)                   # the one shared detector
        email = found[0] if found else ""
        if raw_email.strip() and not email:
            warnings.append(f"{src}: '{raw_email.strip()[:60]}' is not a proper email -- row not imported")
        if not email:
            # A prospect without a proper email is not imported at all: no site, no contact.
            no_email_rows[campaign] = no_email_rows.get(campaign, 0) + 1
            continue
        website = _website(r.get("website", ""), r.get("contact_page", ""))
        lead_id = lead_id_from_url(website) if website else \
            re.sub(r"[^a-z0-9]+", "_", email.split("@")[1].lower()).strip("_")
        domain = domain_from_url(website) if website else email.split("@")[1]
        if not website:
            website = "https://" + domain

        priority, rank = _priority(r.get("priority", ""))
        angle = r.get("suggested_angle", "")
        desc = r.get("description", "")

        lead = {
            "lead_id": lead_id, "campaign_id": campaign,
            "company": r.get("company", ""), "website": website, "domain": domain,
            "country": code, "country_name": COUNTRY_NAMES.get(code, cname),
            "location": r.get("location", ""),
            "prospect_type": r.get("prospect_type", ""),
            "priority": priority,
            "description": desc, "suggested_angle": angle,
            "summary": (desc + "\n\n" + angle).strip() if angle else desc,
            "source_urls": r.get("source_urls", ""),
            "researched_at": r.get("researched_at", ""),
            "notes": r.get("notes", ""),
        }
        if rank is not None:
            lead["priority_rank"] = rank
        pc = _page_count(r.get("page_count", ""))
        if pc:
            lead["page_count"] = pc
        if r.get("sitemap_url"):
            lead["sitemap_url"] = r["sitemap_url"]
        camp = out.setdefault(campaign, {"leads": {}, "contacts": {},
                                         "contacted_skipped": [], "no_email": 0})
        _merge_into(camp["leads"].setdefault(lead_id, {}), lead)

        status = (r.get("prospect_status", "") or "").strip()
        if status and status.lower() != "not contacted":
            camp["contacted_skipped"].append(f"{email} ({status})")
            continue
        did = doc_id_from_email(email)
        contact = {
            "doc_id": did, "email": email, "lead_id": lead_id, "campaign_id": campaign,
            "name": r.get("name", ""), "title": r.get("title", ""),
            "phone": r.get("phone", ""), "website": website, "domain": domain,
            "company": r.get("company", ""),
            "email_type": _email_type(r.get("email_type_raw", ""), email),
            "email_type_raw": r.get("email_type_raw", ""),
            "contact_page": r.get("contact_page", ""),
            "contact_source": r.get("contact_source", ""),
            "notes": r.get("notes", ""),
            "prospect_status": status,
            "prospect_last_contacted": r.get("prospect_last_contacted", ""),
            "prospect_next_followup": r.get("prospect_next_followup", ""),
        }
        _merge_into(camp["contacts"].setdefault(did, {}), contact)

    for campaign, n in no_email_rows.items():
        if campaign in out:
            out[campaign]["no_email"] = n
    if no_email_rows:
        total = sum(no_email_rows.values())
        warnings.append(f"{total} row(s) without a proper email address were NOT imported (no site, no contact)")
    return out, warnings


# ---------------------------------------------------------------------------
# Planning (compare with Firestore)
# ---------------------------------------------------------------------------

def _fetch(db, col, ids: list[str]) -> dict[str, dict]:
    found: dict[str, dict] = {}
    for i in range(0, len(ids), BATCH_SIZE):
        refs = [col.document(x) for x in ids[i:i + BATCH_SIZE]]
        for snap in db.get_all(refs):
            if snap.exists:
                found[snap.id] = snap.to_dict() or {}
    return found


def _diff(old: dict, new: dict) -> dict:
    """{field: (old, new)} for non-empty new values that differ from old."""
    d = {}
    for k, v in new.items():
        if v in ("", None, [], {}):
            continue
        if old.get(k) != v:
            d[k] = (old.get(k), v)
    return d


class Op:
    __slots__ = ("ref", "payload", "kind", "diff", "label")

    def __init__(self, ref, payload, kind, diff, label):
        self.ref, self.payload, self.kind, self.diff, self.label = ref, payload, kind, diff, label


def _plan_docs(col, docs: dict[str, dict], existing: dict[str, dict], *,
               mode: str, new_extra: dict, protected: set = frozenset(),
               repair: dict | None = None) -> tuple[list[Op], dict]:
    """mode 'managed': update changed fields.  'fill': fill only empty fields.
    'create': create missing docs, never touch existing ones.
    repair: {field: default} set on EXISTING docs where the field is missing or empty
    (e.g. campaign_leads.status, which the send filter requires)."""
    ops: list[Op] = []
    stats = {"new": 0, "changed": 0, "unchanged": 0}
    for did, doc in docs.items():
        old = existing.get(did)
        if old is None:
            ops.append(Op(col.document(did), {**doc, **new_extra, "sources": [SOURCE]},
                          "new", {}, did))
            stats["new"] += 1
            continue
        if mode == "create":                 # existing docs are never touched
            stats["unchanged"] += 1
            continue
        if mode == "fill":
            payload = {k: v for k, v in doc.items()
                       if v not in ("", None, [], {}) and old.get(k) in ("", None, [], {})}
            diff = {k: (None, v) for k, v in payload.items()}
        else:
            diff = {k: v for k, v in _diff(old, doc).items() if k not in protected}
            payload = {k: v[1] for k, v in diff.items()}
        for rk, rv in (repair or {}).items():
            if old.get(rk) in ("", None):
                payload[rk] = rv
                diff[rk] = (old.get(rk), rv)
        if diff:
            if SOURCE not in (old.get("sources") or []):
                payload["sources"] = sorted(set(old.get("sources") or []) | {SOURCE})
            ops.append(Op(col.document(did), payload, "changed", diff, did))
            stats["changed"] += 1
        else:
            stats["unchanged"] += 1
    return ops, stats


def _leads_in_other_campaigns(db, campaign_id: str, lead_ids: set) -> int:
    """How many of these sites are also leads in another campaign (information only)."""
    if not lead_ids:
        return 0
    try:
        from google.cloud.firestore_v1.base_query import FieldFilter
        ids, found = list(lead_ids), set()
        for i in range(0, len(ids), 30):
            for d in (db.collection_group(CAMPAIGN_LEADS_SUB)
                      .where(filter=FieldFilter("lead_id", "in", ids[i:i + 30])).stream()):
                parts = d.reference.path.split("/")
                if len(parts) >= 4 and parts[1] != campaign_id:
                    found.add(d.id)
        return len(found)
    except Exception:
        return 0


def build_plan(db, records: dict, *, campaign_only: bool = False,
               force_size: bool = False) -> dict:
    """Compare records with Firestore.  Returns {campaign: {...ops + stats}}."""
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    plan: dict[str, dict] = {}
    try:
        from crm.campaign_import_lib import _contacts_in_other_campaigns
    except Exception:                                    # pragma: no cover
        _contacts_in_other_campaigns = None

    claimed: dict[str, str] = {}     # doc_id -> campaign that takes it in this run
    for campaign, rec in sorted(records.items()):
        camp_ref = db.collection(CAMPAIGNS_COLLECTION).document(campaign)
        camp_exists = camp_ref.get().exists
        leads_col = camp_ref.collection(CAMPAIGN_LEADS_SUB)
        contacts_col = camp_ref.collection(CAMPAIGN_CONTACTS_SUB)

        ex_leads = _fetch(db, leads_col, list(rec["leads"]))
        ex_contacts = _fetch(db, contacts_col, list(rec["contacts"]))

        if not force_size:      # page_count / sitemap_url: only fill empty, never overwrite
            for lid, doc in rec["leads"].items():
                for f in ("page_count", "sitemap_url"):
                    if ex_leads.get(lid, {}).get(f) not in ("", None):
                        doc.pop(f, None)
        lead_ops, lead_stats = _plan_docs(
            leads_col, rec["leads"], ex_leads, mode="managed",
            new_extra={"status": "pending", "synced_at": now},
            repair={"status": "pending"})   # send filter skips leads without a status
        # Duplicate protection for NEW contacts:
        #  1. already active in another campaign in Firestore -> skipped ("reserved")
        #  2. already claimed by another campaign planned in this same run -> skipped
        #  If the Firestore check cannot run, the plan is flagged and --apply refuses.
        new_ids = {d for d in rec["contacts"] if d not in ex_contacts}
        reserved: set = set()
        dup_check_error = ""
        if new_ids:
            if _contacts_in_other_campaigns is None:
                dup_check_error = "duplicate check unavailable (crm.campaign_import_lib)"
            else:
                try:
                    reserved = _contacts_in_other_campaigns(db, campaign, new_ids)
                except Exception as exc:
                    dup_check_error = f"{type(exc).__name__}: {exc}"
        in_run = {d: claimed[d] for d in new_ids if d in claimed}
        contacts = {d: c for d, c in rec["contacts"].items()
                    if d not in reserved and d not in in_run}
        for d in contacts:
            claimed.setdefault(d, campaign)
        leads_elsewhere = _leads_in_other_campaigns(db, campaign, set(rec["leads"]))
        contact_ops, contact_stats = _plan_docs(
            contacts_col, contacts, ex_contacts, mode="managed",
            new_extra={"status": "pending", "sent_at": None, "last_action": "",
                       "last_action_status": "", "created_at": now, "added_at": now},
            protected=PROTECTED_CONTACT_FIELDS)

        entry = {
            "campaign_exists": camp_exists, "camp_ref": camp_ref,
            "lead_ops": lead_ops, "lead_stats": lead_stats,
            "contact_ops": contact_ops, "contact_stats": contact_stats,
            "reserved": sorted(reserved), "no_email": rec["no_email"],
            "in_run_dupes": sorted(f"{d} (already in {c})" for d, c in in_run.items()),
            "dup_check_error": dup_check_error, "leads_elsewhere": leads_elsewhere,
            "contacted_skipped": rec["contacted_skipped"],
            "site_lead_ops": [], "site_contact_ops": [], "email_contact_ops": [],
            "site_lead_stats": None, "site_contact_stats": None,
            "email_contact_stats": None, "missing_from_catalogue": [],
        }

        # Contacts imported earlier but no longer in the catalogue (reported only)
        try:
            from google.cloud.firestore_v1.base_query import FieldFilter
            for snap in contacts_col.where(
                    filter=FieldFilter("sources", "array_contains", SOURCE)).stream():
                if snap.id not in rec["contacts"]:
                    entry["missing_from_catalogue"].append(snap.id)
        except Exception:
            pass

        if not campaign_only:
            sl_col = db.collection(SITE_LEADS_COLLECTION)
            ec_col = db.collection(EMAIL_CONTACTS_COLLECTION)
            site_leads = {
                lid: {"lead_id": lid, "domain": l["domain"], "website": l["website"],
                      "country": l["country"], "country_name": l["country_name"],
                      "company": l["company"], "title": "", "description": l.get("description", ""),
                      "location": l.get("location", ""), "source_query": SOURCE,
                      "query_category": "prospect", "crawled_at": now,
                      "prospect_type": l.get("prospect_type", ""),
                      "page_count": l.get("page_count"), "sitemap_url": l.get("sitemap_url", "")}
                for lid, l in rec["leads"].items()}
            for d in site_leads.values():
                if d.get("page_count") is None:
                    d.pop("page_count", None)
                if not d.get("sitemap_url"):
                    d.pop("sitemap_url", None)
            ex_sl = _fetch(db, sl_col, list(site_leads))
            # existing site_leads belong to the crawler -> create only, never modify
            entry["site_lead_ops"], entry["site_lead_stats"] = _plan_docs(
                sl_col, site_leads, ex_sl, mode="create", new_extra={})
            site_contacts: dict[tuple[str, str], dict] = {}
            for did, c in contacts.items():
                cid = contact_id_from_email(c["email"])
                site_contacts[(c["lead_id"], cid)] = {
                    "contact_id": cid, "email": c["email"], "name": c.get("name", ""),
                    "title": c.get("title", ""), "phone": c.get("phone", ""),
                    "lead_id": c["lead_id"], "website": c["website"],
                    "domain": domain_from_url(c["website"]),
                    "country": rec["leads"][c["lead_id"]]["country"],
                    "country_name": rec["leads"][c["lead_id"]]["country_name"],
                    "found_on": c.get("contact_source") or c.get("contact_page", "")}
            sc_ops: list[Op] = []
            sc_stats = {"new": 0, "changed": 0, "unchanged": 0}
            by_lead: dict[str, dict] = defaultdict(dict)
            for (lid, cid), d in site_contacts.items():
                by_lead[lid][cid] = d
            for lid, docs in by_lead.items():
                col = sl_col.document(lid).collection(SITE_CONTACTS_SUB)
                ops, st = _plan_docs(col, docs, _fetch(db, col, list(docs)),
                                     mode="fill", new_extra={})
                sc_ops += ops
                for k in sc_stats:
                    sc_stats[k] += st[k]
            entry["site_contact_ops"], entry["site_contact_stats"] = sc_ops, sc_stats

            email_docs = {
                did: {"doc_id": did, "email": c["email"], "name": c.get("name", ""),
                      "title": c.get("title", ""), "phone": c.get("phone", ""),
                      "company": c.get("company", ""), "website": c["website"],
                      "domain": domain_from_url(c["website"]),
                      "country": rec["leads"][c["lead_id"]]["country"],
                      "location": rec["leads"][c["lead_id"]].get("location", ""),
                      "email_type": c.get("email_type", ""), "lead_id_site": c["lead_id"],
                      "contact_id": contact_id_from_email(c["email"]), "approved": ""}
                for did, c in contacts.items()}
            ex_ec = _fetch(db, ec_col, list(email_docs))
            entry["email_contact_ops"], entry["email_contact_stats"] = _plan_docs(
                ec_col, email_docs, ex_ec, mode="fill",
                new_extra={"status": "pending", "created_at": now, "campaign": campaign})
        plan[campaign] = entry
    return plan


# ---------------------------------------------------------------------------
# Applying
# ---------------------------------------------------------------------------

def _refresh_counts(db, camp_ref) -> dict:
    """Recount contacts per site: campaign_leads.contact_count (Sites list) plus the
    campaign's contact_count / lead_count.  Contacts and sites stay linked by lead_id."""
    per_lead: dict[str, int] = {}
    total = 0
    for snap in camp_ref.collection(CAMPAIGN_CONTACTS_SUB).select(["lead_id"]).stream():
        total += 1
        lid = (snap.to_dict() or {}).get("lead_id", "")
        per_lead[lid] = per_lead.get(lid, 0) + 1
    batch, pending, leads, lead_ids = db.batch(), 0, 0, set()
    for snap in camp_ref.collection(CAMPAIGN_LEADS_SUB).select(["contact_count"]).stream():
        leads += 1
        lead_ids.add(snap.id)
        want = per_lead.get(snap.id, 0)
        if (snap.to_dict() or {}).get("contact_count") != want:
            batch.update(snap.reference, {"contact_count": want})
            pending += 1
            if pending >= 400:
                batch.commit()
                batch, pending = db.batch(), 0
    if pending:
        batch.commit()
    orphans = sum(n for lid, n in per_lead.items() if lid not in lead_ids)
    if orphans:
        print(f"WARN {camp_ref.id}: {orphans} contact(s) point at a site that is not in the campaign "
              f"-- run 'Update info' on the campaign to re-link them", flush=True)
    return {"contact_count": total, "lead_count": leads}


def apply_plan(db, plan: dict, *, source_files: list[str]) -> dict:
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")

    def commit(ops: list[Op]) -> int:
        n = 0
        for i in range(0, len(ops), BATCH_SIZE):
            batch = db.batch()
            for op in ops[i:i + BATCH_SIZE]:
                batch.set(op.ref, op.payload, merge=True)
            batch.commit()
            n += len(ops[i:i + BATCH_SIZE])
        return n

    written = defaultdict(int)
    for campaign, e in plan.items():
        if not e["campaign_exists"]:
            e["camp_ref"].set({
                "campaign_id": campaign, "status": "draft", "sent_at": None,
                "require_send_confirmation": True,
                "outreach_email_account": "",
                "mail": {"subject": "", "body": "", "type": "plain"},
                "countries": sorted({op.payload.get("country") for op in e["lead_ops"]
                                     if op.payload.get("country")}),
                "status_breakdown": {}, "select_breakdown": {},
                "tier_breakdown": {}, "outreach_breakdown": {},
                "created_at": now, "updated_at": now, "source": SOURCE,
            })
            written["campaigns"] += 1
        written["campaign_leads"]    += commit(e["lead_ops"])
        written["campaign_contacts"] += commit(e["contact_ops"])
        written["site_leads"]        += commit(e["site_lead_ops"])
        written["site_contacts"]     += commit(e["site_contact_ops"])
        written["email_contacts"]    += commit(e["email_contact_ops"])
        upd = {"updated_at": now, "prospect_import": {
            "last_run": now, "files": source_files}}
        try:
            from crm.site_link_lib import ensure_site_links
            ensure_site_links(db, campaign)          # every contact must belong to a site
            upd.update(_refresh_counts(db, e["camp_ref"]))
        except Exception as exc:                     # import itself succeeded
            print(f"WARN could not refresh counts for {campaign}: {exc}", flush=True)
        e["camp_ref"].set(upd, merge=True)
    return dict(written)


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------

def _fmt_stats(s) -> str:
    return "-" if s is None else f"{s['new']:>3} new {s['changed']:>3} chg {s['unchanged']:>3} same"


def print_report(plan: dict, warnings: list[str], *, show_changes: int = 10) -> None:
    for w in warnings:
        print(f"  WARN {w}")
    print()
    for campaign, e in plan.items():
        tag = "(exists)" if e["campaign_exists"] else "(NEW campaign)"
        print(f"== {campaign} {tag}")
        print(f"   campaign_leads    {_fmt_stats(e['lead_stats'])}")
        print(f"   campaign_contacts {_fmt_stats(e['contact_stats'])}")
        if e["site_lead_stats"] is not None:
            print(f"   site_leads        {_fmt_stats(e['site_lead_stats'])}")
            print(f"   site_contacts     {_fmt_stats(e['site_contact_stats'])}")
            print(f"   email_contacts    {_fmt_stats(e['email_contact_stats'])}")
        if e["no_email"]:
            print(f"   {e['no_email']} row(s) without email -> not imported")
        if e["contacted_skipped"]:
            print(f"   {len(e['contacted_skipped'])} contact(s) skipped (catalogue says already contacted): "
                  + ", ".join(e["contacted_skipped"][:5]))
        if e["reserved"]:
            print(f"   {len(e['reserved'])} contact(s) skipped -- active in another campaign: "
                  + ", ".join(e["reserved"][:5]))
        if e["in_run_dupes"]:
            print(f"   {len(e['in_run_dupes'])} contact(s) skipped -- same email in another campaign "
                  f"of this run: " + ", ".join(e["in_run_dupes"][:5]))
        if e["leads_elsewhere"]:
            print(f"   note: {e['leads_elsewhere']} site(s) are also leads in other campaigns "
                  f"(contacts are protected from duplicates, sites are not)")
        if e["dup_check_error"]:
            print(f"   !! duplicate check against other campaigns FAILED: {e['dup_check_error']}")
        if e["missing_from_catalogue"]:
            print(f"   {len(e['missing_from_catalogue'])} contact(s) in Firestore but no longer in the catalogue "
                  f"(kept): " + ", ".join(e["missing_from_catalogue"][:5]))
        shown = 0
        for label, ops in (("lead", e["lead_ops"]), ("contact", e["contact_ops"])):
            for op in ops:
                if op.kind == "changed" and shown < show_changes:
                    parts = [f"{k}: {str(a)[:40]!r} -> {str(b)[:40]!r}"
                             for k, (a, b) in op.diff.items()][:4]
                    print(f"   ~ {label} {op.label}: " + "; ".join(parts))
                    shown += 1
        print()


# ---------------------------------------------------------------------------
# Web import (campaign-import.html -> /api/crm/prospect-import)
# ---------------------------------------------------------------------------

def run_prospect_web_import(db, filename: str, blob: bytes, *, prefix: str = DEFAULT_PREFIX,
                            campaign: str = "", campaign_only: bool = False,
                            dry_run: bool = True) -> dict:
    """Same pipeline as app/prospects_import.py for one uploaded workbook.
    Returns the summary dict the import page shows.  Never writes when dry_run,
    and refuses to write when the duplicate check against other campaigns failed."""
    rows, warnings = parse_catalogue([(filename, blob)])
    records, w2 = build_records(rows, prefix=(prefix or DEFAULT_PREFIX).strip(),
                                only_campaign=(campaign or "").strip())
    warnings += w2
    if not records:
        return {"format": "prospect", "dry_run": dry_run, "rows_parsed": len(rows),
                "campaigns": [], "leads_new": 0, "leads_updated": 0, "contacts_new": 0,
                "contacts_updated": 0, "skipped": 0, "warnings": warnings or
                ["No prospect rows found (need a sheet with Company and Email headers)."]}
    plan = build_plan(db, records, campaign_only=campaign_only)

    failed = [c for c, e in plan.items() if e["dup_check_error"]]
    if failed and not dry_run:
        raise ValueError("Duplicate check against other campaigns failed for "
                         + ", ".join(failed) + " -- nothing was written. Run a dry run for details.")

    camps, tot = [], {"leads_new": 0, "leads_updated": 0, "contacts_new": 0,
                      "contacts_updated": 0, "skipped": 0}
    for cid, e in plan.items():
        skipped = (len(e["reserved"]) + len(e["in_run_dupes"]) + len(e["contacted_skipped"]))
        camps.append({
            "campaign_id": cid, "campaign_exists": e["campaign_exists"],
            "leads": e["lead_stats"], "contacts": e["contact_stats"],
            "site_leads": e["site_lead_stats"], "site_contacts": e["site_contact_stats"],
            "email_contacts": e["email_contact_stats"],
            "no_email": e["no_email"],
            "skipped_other_campaign": e["reserved"][:50],
            "skipped_same_run": e["in_run_dupes"][:50],
            "skipped_contacted": e["contacted_skipped"][:50],
            "dup_check_error": e["dup_check_error"],
            "leads_elsewhere": e["leads_elsewhere"],
            "missing_from_catalogue": len(e["missing_from_catalogue"]),
        })
        tot["leads_new"] += e["lead_stats"]["new"]
        tot["leads_updated"] += e["lead_stats"]["changed"]
        tot["contacts_new"] += e["contact_stats"]["new"]
        tot["contacts_updated"] += e["contact_stats"]["changed"]
        tot["skipped"] += skipped
        if e["dup_check_error"]:
            warnings.append(f"{cid}: duplicate check against other campaigns FAILED "
                            f"({e['dup_check_error']}) -- writing is blocked until it works")

    written = {}
    if not dry_run:
        written = apply_plan(db, plan, source_files=[filename])
    return {"format": "prospect", "dry_run": dry_run, "rows_parsed": len(rows),
            "campaign_id": ", ".join(plan), "campaigns": camps, "written": written,
            "warnings": warnings, **tot}
