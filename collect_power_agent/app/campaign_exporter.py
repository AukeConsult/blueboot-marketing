"""campaign_exporter.py -- Export a campaign to a local Excel file.

Uses the shared campaign_export_data layer, so columns and row structure are
identical to the Google Sheets export triggered by the "Full override" button
in the CRM UI (campaign_export_lib.py).

Output: one .xlsx file with two sheets:
  Summary        -- campaign metadata + country / platform breakdowns
  Leads+Contacts -- one row per contact, lead fields repeated per row

Usage:
    python app/campaign_exporter.py NO_tech_jul01
    python app/campaign_exporter.py NO_tech_jul01 --output /path/to/dir
    python app/campaign_exporter.py --list
"""
from __future__ import annotations

import argparse
from pathlib import Path

import _pathsetup  # noqa: F401  -- sets Windows selector loop + sys.path

_CRM_DIR = str(Path(__file__).resolve().parent.parent / "functions-crm")
import sys as _sys
if _CRM_DIR not in _sys.path:
    _sys.path.insert(0, _CRM_DIR)

from crm.campaign_export_data import (  # noqa: E402
    COLUMNS,
    WRAP_COLS,
    build_summary_rows,
    load_campaign_meta,
    load_rows,
)


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {
        "header_font":  Font(name="Arial", bold=True, color="FFFFFF", size=10),
        "header_fill":  PatternFill("solid", start_color="2D5C8E"),
        "header_align": Alignment(horizontal="center", vertical="center"),
        "cell_font":    Font(name="Arial", size=9),
        "label_font":   Font(name="Arial", bold=True, size=9),
        "wrap":         Alignment(vertical="top", wrap_text=True),
        "nowrap":       Alignment(vertical="top", wrap_text=False),
        "fill_white":   PatternFill("solid", start_color="FFFFFF"),
        "fill_light":   PatternFill("solid", start_color="F2F7FC"),
        "border":       Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        ),
    }


# ---------------------------------------------------------------------------
# Summary sheet
# ---------------------------------------------------------------------------

def _build_summary(ws, campaign_id: str, campaign: dict,
                   rows: list[dict], styles: dict) -> None:
    from openpyxl.styles import Alignment

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34

    def hdr(r: int, label: str) -> None:
        c = ws.cell(r, 1, label)
        c.font      = styles["header_font"]
        c.fill      = styles["header_fill"]
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)

    summary_data = build_summary_rows(campaign_id, campaign, rows)

    # render: blank list = section gap, two-item list = key/value,
    # first row after a header-looking row = sub-header
    r = 1
    in_section = False
    for pair in summary_data:
        if not pair:
            r += 1
            in_section = False
            continue
        if len(pair) == 2:
            key, val = pair
            if not in_section and isinstance(val, str) and val in ("Contacts", "count", ""):
                # treat as section header
                hdr(r, key)
                in_section = True
            else:
                a = ws.cell(r, 1, key)
                b = ws.cell(r, 2, val)
                a.font = styles["label_font"]
                b.font = styles["cell_font"]
        r += 1


# ---------------------------------------------------------------------------
# Leads+Contacts sheet
# ---------------------------------------------------------------------------

def _build_main_sheet(ws, rows: list[dict], styles: dict) -> None:
    from openpyxl.utils import get_column_letter

    # Header row
    for col_idx, (_, label, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(1, col_idx, label)
        c.font      = styles["header_font"]
        c.fill      = styles["header_fill"]
        c.alignment = styles["header_align"]
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20

    # Column widths
    for col_idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        fill = styles["fill_light"] if row_idx % 2 == 0 else styles["fill_white"]
        has_long = False
        for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            value = row.get(key, "")
            c = ws.cell(row_idx, col_idx, value)
            c.fill   = fill
            c.font   = styles["cell_font"]
            c.border = styles["border"]
            if key in WRAP_COLS:
                c.alignment = styles["wrap"]
                if value:
                    has_long = True
            else:
                c.alignment = styles["nowrap"]
        if has_long:
            ws.row_dimensions[row_idx].height = 45


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_campaign(campaign_id: str, output_dir: str | None = None) -> dict:
    from openpyxl import Workbook

    def _get_db():
        try:
            from app.firestore_client import get_firestore
        except ImportError:
            from firestore_client import get_firestore
        return get_firestore()

    db       = _get_db()
    campaign = load_campaign_meta(db, campaign_id)
    rows     = load_rows(db, campaign_id)

    out = Path(output_dir) if output_dir else Path("output") / campaign_id
    out.mkdir(parents=True, exist_ok=True)

    styles = _styles()
    wb     = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _build_summary(ws_summary, campaign_id, campaign, rows, styles)

    ws_main = wb.create_sheet("Leads+Contacts")
    _build_main_sheet(ws_main, rows, styles)

    xlsx_file = out / "campaign.xlsx"
    wb.save(xlsx_file)

    lead_ids = {r["lead_id"] for r in rows}
    print(f"[campaign_export] {campaign_id} → "
          f"{len(lead_ids)} leads / {len(rows)} contacts → {xlsx_file}")

    return {
        "campaign_id":   campaign_id,
        "lead_count":    len(lead_ids),
        "contact_count": len(rows),
        "xlsx_file":     str(xlsx_file),
    }


def _list_campaigns(db) -> list[str]:
    from google.cloud.firestore_v1.base_query import FieldFilter
    docs = (
        db.collection("campaigns")
          .where(filter=FieldFilter("status", "!=", "deleting"))
          .select(["campaign_id"])
          .stream()
    )
    ids = []
    for doc in docs:
        cid = (doc.to_dict() or {}).get("campaign_id", "").strip() or doc.id
        if cid:
            ids.append(cid)
    return sorted(ids)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> None:
    try:
        from dotenv import load_dotenv
        load_dotenv()
    except ImportError:
        pass

    p = argparse.ArgumentParser(
        description="Export a campaign to local Excel (same columns as Google Sheets export)")
    p.add_argument(
        "campaign_id", nargs="?", default=None,
        metavar="CAMPAIGN_ID",
        help="Campaign ID (e.g. NO_tech_jul01). Use --list to see all.")
    p.add_argument(
        "--list", action="store_true",
        help="List available campaign IDs and exit")
    p.add_argument(
        "--output", default=None, metavar="DIR",
        help="Output directory (default: output/<campaign_id>/)")

    args = p.parse_args(argv)

    if args.list:
        def _get_db():
            try:
                from app.firestore_client import get_firestore
            except ImportError:
                from firestore_client import get_firestore
            return get_firestore()
        ids = _list_campaigns(_get_db())
        if ids:
            print("Available campaigns:")
            for cid in ids:
                print(f"  {cid}")
        else:
            print("No campaigns found.")
        return

    if not args.campaign_id:
        p.error("campaign_id is required (or use --list)")

    try:
        result = export_campaign(args.campaign_id, output_dir=args.output)
    except ValueError as exc:
        print(f"  Error: {exc}")
        raise SystemExit(1)

    print(f"\n  Done → {result['xlsx_file']}")
    print(f"  Leads: {result['lead_count']}  Contacts: {result['contact_count']}")


if __name__ == "__main__":
    main()
